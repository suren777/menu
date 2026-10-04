"""CoFID reference data: one-off bulk import into the common ref tables.

Downloads Public Health England's McCance and Widdowson's Composition
of Foods Integrated Dataset 2021 once into .cache/sources/cofid/ and
loads the slices the nutrition work needs into ref_food /
ref_nutrient, tagged with a food_source row and a nutrient_map. A
bulk file, not a crawl, so no politeness rules apply — and nothing
here runs at ingest time: run `uv run menu-cofid-import` to (re)build
the tables.
"""

import logging
from pathlib import Path
from typing import TYPE_CHECKING, Any

import requests
from openpyxl import load_workbook
from sqlalchemy import delete, insert, select

from menu.db.connection import get_session
from menu.db.database import (
    FoodSource,
    NutrientMap,
    RefFood,
    RefNutrient,
    RefPortion,
    initialise,
)

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet
    from sqlalchemy.orm import Session

logger = logging.getLogger(__name__)

CACHE_DIR = Path(".cache/sources/cofid")
SOURCE_NAME = "cofid"
SOURCE_VERSION = "CoFID 2021"

COFID_URL = (
    "https://assets.publishing.service.gov.uk/media/60538b91e90e07527df82ae4/"
    "McCance_Widdowsons_Composition_of_Foods_Integrated_Dataset_2021..xlsx"
)

# The sheets carrying what we slice: edible proportion, the proximates
# (with KCALS) and the one inorganic we keep. Each repeats the food
# rows with its own nutrient columns; columns are located by the
# row-2 nutrient code (the Factors sheet's "Edible proportion" only
# exists as a row-1 header).
SHEETS = ("1.2 Factors", "1.3 Proximates", "1.4 Inorganics")

# Fixed nutrient set, matching FDC's vocabulary so derived values are
# comparable. CoFID's carbohydrate is *available* carbohydrate as
# monosaccharide equivalent — not FDC's by-difference figure — and its
# fibre is AOAC (the older ENGFIB/NSP columns are ignored), so the
# definition each code means is recorded. SATFAC is per 100 g fatty
# acid; SATFOD is the per-100 g-food figure.
NUTRIENT_MAP: list[tuple[str, str, str]] = [
    # (source code, our nutrient, definition)
    ("KCALS", "energy_kcal", "kcal"),
    ("PROT", "protein_g", "by_weight"),
    ("FAT", "fat_g", "by_weight"),
    ("SATFOD", "saturated_fat_g", "by_weight"),
    ("CHO", "carbohydrate_g", "available_monosaccharide"),
    ("TOTSUG", "sugars_g", "total"),
    ("AOACFIB", "fibre_g", "total_dietary"),
    ("NA", "sodium_mg", "by_weight"),
]

CHUNK = 10_000


def _download(url: str) -> Path:
    target = CACHE_DIR / url.rsplit("/", 1)[-1]
    if not target.exists():
        logger.info("Downloading %s", url)
        CACHE_DIR.mkdir(parents=True, exist_ok=True)
        with requests.get(url, stream=True, timeout=300) as response:
            response.raise_for_status()
            with target.open("wb") as handle:
                for chunk in response.iter_content(chunk_size=1 << 20):
                    handle.write(chunk)
    return target


def _amount(value: Any) -> float | None:
    """A per-100 g value from a workbook cell: "Tr" is trace (0.0);
    "N" (not measured), None and empty carry no row."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = value.strip()
    if not text or text == "N":
        return None
    if text == "Tr":
        return 0.0
    return float(text)


def _columns(sheet: Worksheet, wanted: set[str], row: int) -> dict[str, int]:
    """Header key -> column, located by the given header row."""
    header = next(sheet.iter_rows(min_row=row, max_row=row, values_only=True))
    return {
        key: i for i, key in enumerate(header) if isinstance(key, str) and key in wanted
    }


def _rows_by_code(sheet: Worksheet) -> dict[str, tuple[Any, ...]]:
    """Food code -> data row (data starts at row 4). CoFID 2021
    prints one code (13-669) on two rows; last wins."""
    rows: dict[str, tuple[Any, ...]] = {}
    for row in sheet.iter_rows(min_row=4, values_only=True):
        if row[0] is not None:
            rows[str(row[0])] = row
    return rows


def _food_source(session: Session) -> FoodSource:
    source = session.scalar(select(FoodSource).where(FoodSource.name == SOURCE_NAME))
    if source is None:
        source = FoodSource(
            name=SOURCE_NAME,
            version=SOURCE_VERSION,
            licence="Open Government Licence v3.0",
            citation=(
                "Public Health England (2021). McCance and Widdowson's "
                "The Composition of Foods Integrated Dataset 2021. "
                "Contains public sector information licensed under the "
                "Open Government Licence v3.0."
            ),
        )
        session.add(source)
        session.flush()
    return source


def import_cofid(path: Path, session: Session) -> dict[str, int]:
    """Replace the CoFID rows in the common ref tables with the 2021
    workbook's slices. Returns the row counts."""
    source = _food_source(session)
    session.execute(delete(NutrientMap).where(NutrientMap.source_id == source.id))
    session.execute(
        insert(NutrientMap),
        [
            {
                "source_id": source.id,
                "source_nutrient_code": code,
                "nutrient": nutrient,
                "definition": definition,
            }
            for code, nutrient, definition in NUTRIENT_MAP
        ],
    )
    by_code = {
        code: (nutrient, definition)
        for code, nutrient, definition in NUTRIENT_MAP
    }

    ref_food_ids = select(RefFood.id).where(RefFood.source_id == source.id)
    session.execute(
        delete(RefNutrient).where(RefNutrient.ref_food_id.in_(ref_food_ids))
    )
    session.execute(delete(RefPortion).where(RefPortion.ref_food_id.in_(ref_food_ids)))
    session.execute(delete(RefFood).where(RefFood.source_id == source.id))

    foods: list[dict[str, Any]] = []
    nutrients: list[dict[str, Any]] = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        factors = workbook[SHEETS[0]]
        proximates = workbook[SHEETS[1]]
        inorganics = workbook[SHEETS[2]]
        edible_col = _columns(factors, {"Edible proportion"}, row=1)[
            "Edible proportion"
        ]
        factor_rows = _rows_by_code(factors)
        inorganic_rows = _rows_by_code(inorganics)
        proximate_rows = _rows_by_code(proximates)
        # Nutrient code -> the sheet holding it and its column there.
        columns: dict[str, tuple[dict[str, tuple[Any, ...]], int]] = {}
        for sheet, rows in (
            (proximates, proximate_rows),
            (inorganics, inorganic_rows),
        ):
            for source_code, col in _columns(sheet, set(by_code), row=2).items():
                columns[source_code] = (rows, col)

        for code, row in proximate_rows.items():
            ep = _amount(factor_rows[code][edible_col])
            foods.append(
                {
                    "source_id": source.id,
                    "source_food_id": code,
                    "description": row[1],
                    "category": row[3],
                    "refuse_pct": (
                        round((1 - ep) * 100, 2) if ep is not None else None
                    ),
                }
            )
            for source_code, (sheet_rows, col) in columns.items():
                value = _amount(sheet_rows[code][col])
                if value is None:
                    continue
                nutrient, definition = by_code[source_code]
                nutrients.append(
                    {
                        "source_food_id": code,
                        "nutrient": nutrient,
                        "amount_per_100g": value,
                        "definition": definition,
                    }
                )
    finally:
        workbook.close()

    session.execute(insert(RefFood), foods)
    ref_ids = dict(
        session.execute(
            select(RefFood.source_food_id, RefFood.id).where(
                RefFood.source_id == source.id
            )
        ).all()
    )
    nutrient_rows: list[dict[str, Any]] = [
        {"ref_food_id": ref_ids[row.pop("source_food_id")], **row}
        for row in nutrients
    ]

    for start in range(0, len(nutrient_rows), CHUNK):
        session.execute(insert(RefNutrient), nutrient_rows[start : start + CHUNK])
    return {
        "foods": len(foods),
        "portions": 0,  # CoFID publishes no household portions.
        "nutrients": len(nutrients),
    }


def main() -> None:
    logging.basicConfig()
    initialise()
    path = _download(COFID_URL)
    with get_session() as session:
        counts = import_cofid(path, session)
    print(
        f"CoFID imported: {counts['foods']} foods, "
        f"{counts['portions']} portions, {counts['nutrients']} nutrients"
    )


if __name__ == "__main__":
    main()
